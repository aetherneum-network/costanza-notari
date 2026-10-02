"""v2.3 deadline classifier - offline, sockets blocked, fake client. No test calls any API.

OFF by default; refused without a key; the effort knob; the prompt that is sent; the deterministic gate
(every refusal path); the 12 D14 answers whose evidence joined fragments with '...' abstain; the end-to-end
path through the pipeline (commit, refusal, transport error, configuration fault, caller's stop).
"""
import contextlib
import hashlib
import io
import json
import os
import socket
import unittest
from unittest import mock

from tests._util import AS_OF, ROOT, ensure_corpus, main_run, run_pipeline
from pipeline import classifier_gate as gate, classify, llm_classifier as L, run as runner
from pipeline.lib import jsonio, tzrome

CONFIG = jsonio.read(ROOT / "corpus" / "config.json")
FIXTURE = ROOT / "tests" / "fixtures" / "classifier_d14_joined.json"
USAGE = {"input_tokens": 1200, "cache_creation_input_tokens": 0, "cache_read_input_tokens": 9000,
         "output_tokens": 700}
_REAL = {}


def setUpModule():
    def blocked(*_a, **_k):
        raise OSError("network blocked in the offline test-suite")
    _REAL["create_connection"], _REAL["getaddrinfo"] = socket.create_connection, socket.getaddrinfo
    _REAL["connect"] = socket.socket.connect
    socket.create_connection = socket.getaddrinfo = blocked
    socket.socket.connect = blocked


def tearDownModule():
    socket.create_connection, socket.getaddrinfo = _REAL["create_connection"], _REAL["getaddrinfo"]
    socket.socket.connect = _REAL["connect"]


# ------------------------------------------------------------------------------------------- fake client
class _Block:
    def __init__(self, text):
        self.type, self.text = "text", text


class _Usage(dict):
    def model_dump(self, mode="json", exclude_none=True):
        return dict(self)


class _Msg:
    def __init__(self, answer=None, stop_reason="end_turn", text=None):
        self.stop_reason = stop_reason
        body = text if text is not None else (json.dumps(answer) if answer is not None else None)
        self.content = [] if body is None else [_Block(body)]
        self.usage = _Usage(USAGE)        # like the SDK: a streamed message has no _request_id


class _Stream:
    request_id = "req_fake"                # like the SDK: the id is a header of the stream's response

    def __init__(self, result):
        self.result = result

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False

    def get_final_message(self):
        if isinstance(self.result, BaseException):
            raise self.result
        return self.result


class FakeClient:
    """Answers by record id (first line of the user message); records every request it receives."""

    def __init__(self, by_record=None, default=None):
        self.by_record, self.default, self.calls = by_record or {}, default, []
        self.messages = self

    def stream(self, **params):
        self.calls.append(params)
        rid = params["messages"][0]["content"].split("\n", 1)[0].removeprefix("RECORD ")
        res = self.by_record.get(rid, self.default)
        return _Stream(res(params) if callable(res) else res)


def answer(deadline="RECUPERARE", nature="unknown", act_type="RECUPERARE", evidence="", computation="",
           confidence="high"):
    return {"deadline": deadline, "nature": nature, "act_type": act_type, "evidence": evidence,
            "computation": computation, "confidence": confidence}


ABSTAIN = _Msg(answer())


class APIConnectionError(Exception):
    """Same class name as the SDK's: the classifier sorts errors by class name, without importing the SDK."""


class BadRequestError(Exception):
    status_code = 400


def ctx():
    return classify.build_context(CONFIG, tzrome.parse_iso(AS_OF))


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def quiet(fn, *a, **k):
    with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
        return fn(*a, **k)


# ------------------------------------------------------------------------------------------- OFF / key
class OffByDefault(unittest.TestCase):
    def test_parser_and_factory_default_to_off(self):
        args = runner.parser().parse_args([])
        self.assertIsNone(args.llm)
        self.assertIsNone(args.llm_effort)
        for name in (None, "", "none", "disabled"):
            self.assertFalse(L.get(name).enabled, name)

    def test_cached_main_run_has_no_classifier_anywhere(self):
        base = main_run()
        st = jsonio.read(base / "work" / "state" / "05_classification.json")
        self.assertEqual(st["llm"], {"enabled": False})
        self.assertFalse(any("classifier" in r for r in st["records"]))
        led = jsonio.read(base / "work" / "state" / "07_ledger.json")
        self.assertFalse(any("classifier" in u["facts"] for u in led["view"]))
        rep = jsonio.read(base / "work" / "run_report.json")
        self.assertEqual(rep["stages"]["s5_classify"]["llm"], False)
        self.assertNotIn("classifier", rep["stages"]["s5_classify"])

    def test_key_and_effort_in_the_environment_do_not_switch_it_on(self):
        env = {"ANTHROPIC_API_KEY": "test-placeholder-not-a-key", L.EFFORT_ENV: "high"}
        with mock.patch.dict(os.environ, env), \
                mock.patch.object(L, "AnthropicLLM", side_effect=AssertionError("classifier built while OFF")):
            code, base = quiet(run_pipeline, "v23-off-env")
        self.assertEqual(code, 0)
        st = jsonio.read(base / "work" / "state" / "05_classification.json")
        self.assertFalse(st["llm"]["enabled"])
        self.assertFalse(any("classifier" in r for r in st["records"]))
        for name in ("master_index.xlsx", "report.docx"):     # OFF output: byte-identical to v2.2.1's path
            self.assertEqual(sha(base / "work" / "build" / name), sha(main_run() / "work" / "build" / name), name)


class RefusedWithoutKey(unittest.TestCase):
    def test_factory_refuses_without_key(self):
        with mock.patch.dict(os.environ, {"ANTHROPIC_API_KEY": "  "}):
            with self.assertRaises(L.ClassifierConfigError) as cm:
                L.get("anthropic")
        self.assertIn("ANTHROPIC_API_KEY", str(cm.exception))

    def test_run_fails_before_any_stage_without_key(self):
        with mock.patch.dict(os.environ, {"ANTHROPIC_API_KEY": ""}):
            code, base = quiet(run_pipeline, "v23-nokey", "--llm", "anthropic")
        self.assertEqual(code, 3)
        rep = jsonio.read(base / "work" / "run_report.json")
        self.assertEqual(rep["status"], "FAILED")
        self.assertIn("ClassifierConfigError", rep["failure"])
        self.assertIn("ANTHROPIC_API_KEY", rep["failure"])
        self.assertEqual(rep["stages"], {})                       # nothing ran
        self.assertFalse((base / "work" / "state" / "01_enumeration.json").exists())
        self.assertIsNone(rep.get("published"))

    def test_unknown_backend_fails(self):
        with self.assertRaises(L.ClassifierConfigError):
            L.get("another-vendor")

    def test_effort_flag_without_llm_is_a_usage_error(self):
        with self.assertRaises(SystemExit):
            quiet(runner.main, ["--llm-effort", "low"])


class EffortKnob(unittest.TestCase):
    def test_default_is_medium(self):
        self.assertEqual(L.DEFAULT_EFFORT, "medium")
        self.assertEqual(L.resolve_effort(None, {}), "medium")

    def test_environment_variable(self):
        for e in ("low", "medium", "high"):
            self.assertEqual(L.resolve_effort(None, {L.EFFORT_ENV: e}), e)
        self.assertEqual(L.resolve_effort(None, {L.EFFORT_ENV: " High "}), "high")

    def test_flag_wins_over_environment(self):
        self.assertEqual(L.resolve_effort("high", {L.EFFORT_ENV: "low"}), "high")

    def test_other_values_are_refused(self):
        for bad in ("xhigh", "max", "fast", "0"):
            with self.assertRaises(L.ClassifierConfigError):
                L.resolve_effort(None, {L.EFFORT_ENV: bad})
            with self.assertRaises(L.ClassifierConfigError):
                L.resolve_effort(bad, {})
        with self.assertRaises(SystemExit):
            quiet(runner.parser().parse_args, ["--llm", "anthropic", "--llm-effort", "xhigh"])

    def test_the_request_carries_the_effort(self):
        rec, env, txt = _fixture_case(0)
        for e in ("low", "medium", "high"):
            llm = L.AnthropicLLM(effort=e, client=FakeClient(default=ABSTAIN))
            self.assertEqual(llm.params(rec, env, txt, tzrome.parse_iso(AS_OF))["output_config"]["effort"], e)
        with mock.patch.dict(os.environ, {L.EFFORT_ENV: "low"}):
            self.assertEqual(L.get("anthropic", client=FakeClient()).effort, "low")
            self.assertEqual(L.get("anthropic", effort="high", client=FakeClient()).effort, "high")


# ------------------------------------------------------------------------------------------- prompt
def _fixture():
    return jsonio.read(FIXTURE)


def _fixture_case(i):
    c = _fixture()["cases"][i]
    return c["rec"], c["env"], c["txt"]


class PromptContent(unittest.TestCase):
    def setUp(self):
        self.rec, self.env, self.txt = _fixture_case(0)
        self.client = FakeClient(default=ABSTAIN)
        self.llm = L.AnthropicLLM(effort="medium", client=self.client)
        self.out = self.llm.propose(self.rec, self.env, self.txt, tzrome.parse_iso(AS_OF))
        self.params = self.client.calls[0]

    def test_model_streaming_and_caching(self):
        self.assertEqual(self.params["model"], "claude-opus-5-5")
        self.assertEqual(self.params["max_tokens"], L.MAX_TOKENS)
        self.assertEqual(self.params["system"][0]["cache_control"], {"type": "ephemeral"})
        self.assertEqual(self.out["model"], "claude-opus-5-5")

    def test_no_parameter_opus_5_5_refuses_and_no_fallback(self):
        dumped = json.dumps(self.params)
        for k in ("temperature", "top_p", "top_k", "thinking", "tool_choice", "tools", "fallbacks", "betas"):
            self.assertNotIn(k, self.params, k)
        self.assertNotIn("budget_tokens", dumped)

    def test_request_id_comes_from_the_stream(self):
        self.assertEqual(self.out["request_id"], "req_fake")

    def test_strict_json_schema(self):
        fmt = self.params["output_config"]["format"]
        self.assertEqual(fmt, {"type": "json_schema", "schema": L.OUTPUT_SCHEMA})
        self.assertFalse(L.OUTPUT_SCHEMA["additionalProperties"])
        self.assertEqual(set(L.OUTPUT_SCHEMA["required"]), set(L.OUTPUT_SCHEMA["properties"]))

    def test_system_prompt_holds_contract_terms_md_and_table(self):
        system = self.params["system"][0]["text"]
        self.assertIn((ROOT / "docs" / "TERMS.md").read_text(encoding="utf-8"), system)   # verbatim
        terms = jsonio.read(ROOT / "rules" / "terms.json")
        for t in terms["terms"]:
            self.assertIn(f"{t['id']} = {t['doc_type']}", system)
        for d in terms["expects_deadline"]:
            self.assertIn(d, system)
        self.assertIn(f"Reference instant (as_of): {AS_OF}", system)
        self.assertIn("Answering RECUPERARE is always safe", system)
        self.assertIn("ONE contiguous passage", system)
        self.assertIn('no fragments joined by "..."', system)

    def test_user_message_holds_the_record_text(self):
        user = self.params["messages"][0]["content"]
        self.assertTrue(user.startswith(f"RECORD {self.rec['record_id']}\n"))
        dc = self.env["daticert"]
        for v in (dc["data"], dc["mittente"], dc["oggetto"], self.rec["doc_type"], self.rec["sender_class"],
                  self.rec["recuperare_reasons"]["deadline"]):
            self.assertIn(str(v), user)
        principal = next(d for d in self.txt["documents"] if d["name"] == self.txt["principal"])
        self.assertIn(f"DOCUMENT 1 (principal) {principal['name']}:", user)
        self.assertIn(principal["text"].strip()[:400], user)
        self.assertIn((self.env["inner"]["body"] or "").strip()[:200] or "[empty]", user)


# ------------------------------------------------------------------------------------------- gate
class GateRules(unittest.TestCase):
    def test_inline_tests_of_every_rule_pass(self):
        self.assertEqual(gate.run_inline_tests(ctx()), [])

    def test_every_rule_has_a_test_that_it_decides(self):
        rules = gate.load()["rules"]
        self.assertEqual([r["id"] for r in rules], [f"G-{i:02d}" for i in range(1, len(rules) + 1)])
        for r in rules:
            self.assertIn(r["check"], gate.CHECKS, r["id"])
            self.assertTrue(r["rationale"], r["id"])
            self.assertIn(r["id"], [t["expect"]["rule"] for t in r["tests"]], r["id"])
            self.assertIn(r["kind"], ("no_answer", "model_abstained", "refusal"), r["id"])

    def _eval(self, a, rec=None, source=None, status="ok"):
        base = {"dates": [], "deadlines": [], "doc_type": "atto_precetto", "notification_date": "2026-10-05",
                "notification_note": None, "pec_time": "2026-10-05T10:00:00+02:00"}
        return gate.evaluate(status, a, {**base, **(rec or {})}, source or gate.load()["test_source"], ctx())

    def test_each_refusal_path_records_rule_kind_and_reason(self):
        ev = "Udienza fissata per il giorno 14/11/2026"
        cases = [
            ("G-01", None, {"status": "max_tokens"}),
            ("G-02", answer(), {}),
            ("G-03", answer("NONE", "none", "atto_precetto", "ATTO DI PRECETTO"), {}),
            ("G-04", answer("2026-11-14", "conditional", "atto_precetto", ev), {}),
            ("G-05", answer("2026-11-14", "actionable", "atto_precetto", ""), {}),
            ("G-06", answer("2026-11-14", "actionable", "atto_precetto", "ATTO DI PRECETTO ... 14/11/2026"), {}),
            ("G-07", answer("2026-11-14", "actionable", "atto_precetto", "udienza fissata il 14/11/2026"), {}),
            ("G-08", answer("2026-11-15", "actionable", "atto_precetto", ev), {}),
            ("G-09", answer("2026-11-14", "actionable", "atto_precetto", ev),
             {"rec": {"dates": [{"date": "2026-11-14", "nature": "conditional", "rule": "N-X"}]}}),
            ("G-10", answer("2026-10-15", "computed", "atto_precetto", "entro dieci giorni dalla notifica"),
             {"rec": {"notification_date": None, "notification_note": "internal forward"}}),
            ("G-11", answer("2026-10-15", "computed", "decreto_ingiuntivo", "entro dieci giorni dalla notifica"), {}),
            ("G-12", answer("2026-10-04", "computed", "atto_precetto", "entro dieci giorni dalla notifica"), {}),
            ("G-13", answer("2026-10-14", "computed", "atto_precetto", "entro dieci giorni dalla notifica"), {}),
            ("G-14", answer("2026-11-14", "actionable", "atto_precetto", ev),
             {"rec": {"deadlines": [{"date": "2026-10-30", "nature": "computed"}]}}),
            ("G-15", answer("2026-11-14", "actionable", "atto_precetto", ev, confidence="low"), {}),
        ]
        for rid, a, kw in cases:
            with self.subTest(rid):
                got = self._eval(a, kw.get("rec"), status=kw.get("status", "ok"))
                self.assertFalse(got["passed"])
                self.assertEqual(got["rule"], rid, got["reason"])
                self.assertTrue(got["reason"])
                self.assertEqual(got["kind"], {"G-01": "no_answer", "G-02": "model_abstained"}.get(rid, "refusal"))

    def test_a_good_answer_passes(self):
        got = self._eval(answer("2026-11-14", "actionable", "atto_precetto", "Udienza fissata per il giorno 14/11/2026"))
        self.assertTrue(got["passed"], got)
        got = self._eval(answer("2026-10-15", "computed", "atto_precetto", "entro dieci giorni dalla notifica"))
        self.assertTrue(got["passed"], got)

    def test_first_refusal_wins(self):
        a = answer("2026-11-14", "actionable", "atto_precetto", "ATTO ... 14/11/2026", confidence="low")
        self.assertEqual(self._eval(a)["rule"], "G-06")

    def test_verbatim_is_whitespace_insensitive_only(self):
        src = "ATTO DI PRECETTO\nSi intima di pagare entro dieci\n  giorni dalla notifica."
        ok = answer("2026-10-15", "computed", "atto_precetto", "entro dieci giorni dalla notifica")
        self.assertTrue(self._eval(ok, source=src)["passed"])
        case = answer("2026-10-15", "computed", "atto_precetto", "Entro dieci giorni dalla notifica")
        self.assertEqual(self._eval(case, source=src)["rule"], "G-07")

    def test_an_ellipsis_written_in_the_record_itself_is_verbatim(self):
        src = "Udienza fissata ... per il giorno 14/11/2026."
        a = answer("2026-11-14", "actionable", "atto_precetto", "Udienza fissata ... per il giorno 14/11/2026")
        self.assertTrue(self._eval(a, source=src)["passed"])

    def test_non_answers_are_classified(self):
        for msg, want in ((_Msg(stop_reason="refusal"), "refusal"), (_Msg(stop_reason="max_tokens"), "max_tokens"),
                          (_Msg(text=None), "no_text_block"), (_Msg(text="{not json"), "invalid_json"),
                          (_Msg(text=json.dumps({"deadline": "2026-11-14"})), "invalid_schema"),
                          (_Msg({**answer(), "nature": "maybe"}), "invalid_schema"), (ABSTAIN, "ok")):
            with self.subTest(want):
                self.assertEqual(L.parse_answer(msg)[0], want)


class D14JoinedEvidence(unittest.TestCase):
    """D14 (2026-10-02): 12 correct dates rested on fragments joined by '...'. With the gate they abstain."""

    def test_the_twelve_records_abstain_at_every_effort(self):
        cases = _fixture()["cases"]
        self.assertEqual(len(cases), 12)
        c5 = ctx()
        n = 0
        for c in cases:
            for effort, a in sorted(c["d14_answers"].items()):
                with self.subTest(record=c["record_id"], effort=effort):
                    self.assertIn("...", a["evidence"])
                    self.assertEqual(a["deadline"], c["gold_deadline"])       # it was a correct date
                    llm = L.AnthropicLLM(effort=effort, client=FakeClient(default=_Msg(a)))
                    out = llm.propose(c["rec"], c["env"], c["txt"], tzrome.parse_iso(AS_OF))
                    self.assertEqual(out["status"], "ok")
                    g = gate.evaluate(out["status"], out["answer"], c["rec"],
                                      gate.source_text(c["env"], c["txt"]), c5)
                    self.assertFalse(g["passed"])
                    self.assertEqual(g["rule"], "G-06", g["reason"])
                    n += 1
        self.assertEqual(n, 36)

    def test_one_of_their_fragments_alone_would_pass(self):
        """The gate refuses the join, not the record: the fragment that states the term passes and gives the
        same (correct) date, counted by the pipeline itself."""
        c5 = ctx()
        for c in _fixture()["cases"]:
            a = c["d14_answers"]["medium"]
            frag = next(f.strip() for f in a["evidence"].split("...") if "giorni" in f)
            g = gate.evaluate("ok", {**a, "evidence": frag}, c["rec"], gate.source_text(c["env"], c["txt"]), c5)
            self.assertTrue(g["passed"], (c["record_id"], g))


# ------------------------------------------------------------------------------------------- end to end
SIMULATED = "ARC-0084"          # dev corpus: an atto di citazione with a hearing on 12/11/2026


def _abstaining_rules(real):
    """Simulate a rule abstention on one record (the dev corpus has no record the classifier can fill)."""
    def wrapped(env, sig, txt, ctx_):
        rec = real(env, sig, txt, ctx_)
        if rec["record_id"] == SIMULATED:
            rec.update(deadline="RECUPERARE", deadline_nature=None, deadline_status=None, days_left=None,
                       urgency="MAXIMUM", deadlines=[])
            rec["recuperare_fields"] = sorted(set(rec["recuperare_fields"]) | {"deadline"})
            rec["recuperare_reasons"]["deadline"] = "simulated rule abstention (test)"
        return rec
    return wrapped


class EndToEnd(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        ensure_corpus()
        hearing = "a comparire all'udienza del 12/11/2026 innanzi al"
        cls.client = FakeClient(by_record={
            SIMULATED: _Msg(answer("2026-11-12", "actionable", "atto_citazione", hearing)),
            "ARC-0016": _Msg(answer("2026-12-01", "computed", "decreto_ingiuntivo",
                                    "Inoltro decreto ingiuntivo ricevuto a mezzo posta")),
            "ARC-0027": APIConnectionError("connection reset (fake)"),
        }, default=ABSTAIN)
        real_get = L.get
        with mock.patch.object(classify, "classify_record", _abstaining_rules(classify.classify_record)), \
                mock.patch.object(L, "get", lambda name, effort=None, client=None:
                                  real_get(name, effort, client=cls.client)), \
                mock.patch.dict(os.environ, {L.EFFORT_ENV: "low"}):
            cls.code, cls.base = quiet(run_pipeline, "v23-e2e", "--llm", "anthropic", "--llm-effort", "high")
        cls.st = jsonio.read(cls.base / "work" / "state" / "05_classification.json")
        cls.recs = {r["record_id"]: r for r in cls.st["records"]}

    def test_run_ok_and_only_recuperare_deadlines_were_sent(self):
        self.assertEqual(self.code, 0)
        sent = sorted(p["messages"][0]["content"].split("\n", 1)[0][7:] for p in self.client.calls)
        want = sorted(rid for rid, r in self.recs.items() if "classifier" in r)
        self.assertEqual(sent, want)
        self.assertEqual(len(sent), 8)              # the 7 RECUPERARE deadlines of the dev corpus + the simulated one
        self.assertTrue(all(p["output_config"]["effort"] == "high" for p in self.client.calls))  # the flag won

    def test_gate_passed_answer_fills_the_deadline_marked(self):
        r = self.recs[SIMULATED]
        self.assertEqual((r["deadline"], r["deadline_nature"], r["deadline_status"]), ("2026-11-12", "actionable", "open"))
        self.assertNotIn("deadline", r["recuperare_fields"])
        self.assertNotIn("deadline", r["recuperare_reasons"])
        self.assertEqual(r["classifier"]["rules_reason"], "simulated rule abstention (test)")   # moved, not lost
        self.assertTrue(r["classifier"]["gate"]["passed"])
        self.assertEqual(r["deadlines"][-1]["source"], "classifier")
        self.assertIn("classifier claude-opus-5-5 (high)", r["rule_trace"]["deadline"])
        self.assertEqual(r["urgency"], "HIGH")
        self.assertEqual(r["classifier"]["usage"], USAGE)
        cons = {x["record_id"]: x for x in jsonio.read(self.base / "work" / "state" / "06_consolidation.json")["records"]}
        self.assertEqual(cons[SIMULATED]["deadline"], "2026-11-12")
        led = jsonio.read(self.base / "work" / "state" / "07_ledger.json")
        facts = next(u["facts"] for u in led["view"] if u["facts"]["record_id"] == SIMULATED)
        self.assertTrue(facts["classifier"]["gate"]["passed"])

    def test_refused_and_failed_calls_abstain_with_their_reason(self):
        r = self.recs["ARC-0016"]
        self.assertEqual(r["deadline"], "RECUPERARE")
        self.assertEqual((r["classifier"]["gate"]["rule"], r["classifier"]["gate"]["kind"]), ("G-10", "refusal"))
        self.assertIn("deadline", r["recuperare_reasons"])
        t = self.recs["ARC-0027"]
        self.assertEqual(t["deadline"], "RECUPERARE")
        self.assertEqual(t["classifier"]["status"], "transport_error")
        self.assertEqual(t["classifier"]["gate"]["rule"], "G-01")
        s = self.st["llm"]["classifier"]
        self.assertEqual((s["called"], s["committed"], s["gate_refused"], s["no_answer"], s["model_abstained"]),
                         (8, 1, 1, 1, 5))
        self.assertEqual(s["refused_by_rule"], {"G-10": 1})
        rep = jsonio.read(self.base / "work" / "run_report.json")
        self.assertEqual(rep["stages"]["s5_classify"]["classifier"]["committed"], 1)

    def test_workbook_shows_the_classifier(self):
        from openpyxl import load_workbook
        wb = load_workbook(self.base / "work" / "build" / "master_index.xlsx")
        basis = [row for row in wb["Deadlines"].iter_rows(min_row=2, values_only=True) if row[3] == SIMULATED]
        self.assertEqual(basis[0][0], "2026-11-12")
        self.assertIn("classifier claude-opus-5-5 (high), gate passed", basis[0][6])
        idx = [row for row in wb["Index"].iter_rows(min_row=7, values_only=True) if row[19] == SIMULATED]
        self.assertEqual(idx[0][5], "actionable (classifier)")
        rec_rows = {(row[0], row[1]): row[2] for row in wb["RECUPERARE"].iter_rows(min_row=2, values_only=True)}
        self.assertIn("gate G-10 refused it", rec_rows[("ARC-0016", "deadline")])
        self.assertIn("no usable answer (transport_error)", rec_rows[("ARC-0027", "deadline")])
        self.assertNotIn((SIMULATED, "deadline"), rec_rows)


class LoudFailures(unittest.TestCase):
    def _run(self, name, client):
        real_get = L.get
        with mock.patch.object(L, "get", lambda n, effort=None, client_=None, **_: real_get(n, effort, client=client)):
            return quiet(run_pipeline, name, "--llm", "anthropic")

    def test_api_configuration_error_fails_the_run(self):
        code, base = self._run("v23-4xx", FakeClient(default=BadRequestError("bad request (fake)")))
        self.assertEqual(code, 3)
        rep = jsonio.read(base / "work" / "run_report.json")
        self.assertEqual(rep["status"], "FAILED")
        self.assertIn("ClassifierConfigError", rep["failure"])
        self.assertIsNone(rep.get("published"))

    def test_a_callers_stop_is_never_swallowed(self):
        code, base = self._run("v23-stop", FakeClient(default=L.ClassifierStop("spending cap reached (fake)")))
        self.assertEqual(code, 3)
        rep = jsonio.read(base / "work" / "run_report.json")
        self.assertIn("ClassifierStop", rep["failure"])


if __name__ == "__main__":
    unittest.main()
