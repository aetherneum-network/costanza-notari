import datetime as dt
import unittest

from tests._util import main_run, tmpdir
from pipeline import publish as pub
from pipeline.lib import tzrome


class IndexAndReportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.base = main_run()
        from openpyxl import load_workbook
        cls.wb = load_workbook(cls.base / "work" / "build" / "master_index.xlsx")

    def test_header_shows_data_as_of_and_release_banner(self):
        ws = self.wb["Index"]
        self.assertEqual(ws["A2"].value, "Data as of 2026-10-21 09:40 (Europe/Rome, UTC+02:00)")
        self.assertTrue(ws["A3"].value.startswith("RELEASE: OK"))
        self.assertIn("SYNTHETIC DATA", ws["A4"].value)

    def test_recuperare_is_bold_red_on_yellow(self):
        ws = self.wb["Index"]
        cells = [c for row in ws.iter_rows(min_row=7) for c in row if c.value == "RECUPERARE"]
        self.assertTrue(cells)
        for c in cells:
            self.assertTrue(c.font.bold)
            self.assertEqual(c.font.color.rgb[-6:], "FF0000")
            self.assertEqual(c.fill.start_color.rgb[-6:], "FFFF00")

    def test_urgency_colours_and_conditional_formatting(self):
        ws = self.wb["Index"]
        seen = {}
        for row in ws.iter_rows(min_row=7):
            seen[row[3].value] = row[3].fill.start_color.rgb[-6:]
        self.assertEqual(seen.get("MAXIMUM"), "C00000")
        self.assertEqual(seen.get("INFORMATIONAL"), "D9D9D9")
        ranges = [str(r.sqref) for r in ws.conditional_formatting]
        self.assertTrue(any(r.startswith("D7:D") for r in ranges))
        self.assertTrue(any(r.startswith("A7:U") for r in ranges))

    def test_edition_bound_amounts(self):
        ws = self.wb["Index"]
        amounts = [row[8].value for row in ws.iter_rows(min_row=7) if row[8].value not in ("-", "RECUPERARE")]
        self.assertTrue(amounts)
        self.assertTrue(all(" - per notice ref. " in a and " of 20" in a for a in amounts))

    def test_signatures_sheet_keeps_integrity_and_chain_apart(self):
        rows = list(self.wb["Signatures"].iter_rows(min_row=2, values_only=True))
        self.assertTrue(any(r[3] == "ok" and r[4] == "false" for r in rows))

    def test_docx_core_properties_are_fixed(self):
        from docx import Document
        d = Document(str(self.base / "work" / "build" / "report.docx"))
        self.assertEqual(d.core_properties.created, dt.datetime(2026, 10, 21, 7, 40, tzinfo=dt.timezone.utc))
        self.assertEqual(d.core_properties.author, "Costanza Notari (synthetic alumna)")
        text = "\n".join(p.text for p in d.paragraphs)
        self.assertIn("Data as of 2026-10-21 09:40", text)
        self.assertIn("SYNTHETIC", text)

    def test_published_to_store(self):
        v = pub.reader_view(self.base / "store", tzrome.parse_iso("2026-10-21T10:00:00+02:00"))
        self.assertEqual((v["version"], v["stale"]), (1, False))


class PublishTests(unittest.TestCase):
    def test_if_version_conflict_and_supersede_not_delete(self):
        d = tmpdir("s8-publish")
        f = d / "a.txt"
        f.write_text("one", encoding="utf-8")
        r1 = pub.publish(d / "store", {"a.txt": f}, if_version=0, as_of="2026-10-20T09:00:00+02:00", run_id="R1")
        f.write_text("two", encoding="utf-8")
        pub.publish(d / "store", {"a.txt": f}, if_version=r1["version"], as_of="2026-10-21T09:00:00+02:00", run_id="R2")
        self.assertEqual((d / "store" / "_SUPERSEDED" / "v0001" / "a.txt").read_text(encoding="utf-8"), "one")
        with self.assertRaises(pub.PublishConflict):
            pub.publish(d / "store", {"a.txt": f}, if_version=1, as_of="2026-10-21T10:00:00+02:00", run_id="R3")

    def test_reader_flags_stale_data(self):
        d = tmpdir("s8-stale")
        f = d / "a.txt"
        f.write_text("x", encoding="utf-8")
        pub.publish(d / "store", {"a.txt": f}, if_version=0, as_of="2026-10-20T01:05:00+02:00", run_id="R1")
        v = pub.reader_view(d / "store", tzrome.parse_iso("2026-10-21T09:40:00+02:00"))
        self.assertTrue(v["stale"])
        self.assertTrue(v["banner"].startswith("STALE DATA"))


if __name__ == "__main__":
    unittest.main()
