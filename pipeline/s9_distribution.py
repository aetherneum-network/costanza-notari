"""Stage 9 - distribution check: superseded values in OUTGOING documents (L2).

The index is right; are the documents leaving the office right too? For each
entry of ``superseded_values.json`` the scanner looks for the *old* value in
DOCX, XLSX (including the raw numeric form, e.g. ``3415.2``), EML and
Markdown, and applies the precision guards before flagging:

* G1 snapshot by name  - archive / snapshot / _SUPERSEDED files are skipped;
* G2 snapshot by time  - files untouched since ``since`` are skipped;
* G3 comparison        - old and new value side by side: not an error;
* G4 context           - "previously", historical totals, quoted replies;
* G5 other sense       - the segment cites only OTHER dossier references
                         (PR-/EX-/ARC-/SV- ids), names another counterparty, or
                         the document has no link at all to the entry's
                         counterparty / dossier;
* G6 owner closure     - an owner may declare ``already_fixed``.

Findings are grouped by owner; precision is tracked as a metric.
"""
from __future__ import annotations

import datetime as _dt
import email
import email.utils
import email.policy
import os
import re
from pathlib import Path

from . import entities as ent
from .amounts import format_amount, format_amount_it
from .lib import jsonio
from .rules_engine import RULES_DIR


MONTHS = ["gennaio", "febbraio", "marzo", "aprile", "maggio", "giugno", "luglio", "agosto", "settembre",
          "ottobre", "novembre", "dicembre"]


def _is_date(v: str) -> bool:
    return bool(re.fullmatch(r"\d{4}-\d{2}-\d{2}", v or ""))


def text_forms(v: str) -> list[str]:
    if _is_date(v):
        d = _dt.date.fromisoformat(v)
        return sorted({v, d.strftime("%d/%m/%Y"), d.strftime("%d.%m.%Y"), f"{d.day} {MONTHS[d.month - 1]} {d.year}",
                       f"{d.day:02d} {MONTHS[d.month - 1]} {d.year}"}, key=len, reverse=True)
    en, it = format_amount(v), format_amount_it(v)
    plain = v
    plain_it = v.replace(".", ",")
    short = v.rstrip("0").rstrip(".") if "." in v else v
    forms = [en, it, plain, plain_it]
    if short not in forms:
        forms.append(short)
    return sorted(set(forms), key=len, reverse=True)


def form_regex(v: str) -> re.Pattern:
    alts = "|".join(re.escape(f) for f in text_forms(v))
    return re.compile(r"(?<![\d.,])(?:" + alts + r")(?![\d]|[.,]\d)")


REF_RX = re.compile(r"\b(PR-[A-Z0-9-]+[A-Z0-9]|EX-\d+|ARC-\d{4}|SV-[0-9A-Z]+(?:-[a-z_]+)?|B-[0-9a-f]{12})\b")


def own_refs(e: dict) -> set:
    p = e.get("proof", {})
    return {x for x in (e.get("id"), e.get("base_id"), e.get("pratica"), p.get("document"), p.get("supersedes"),
                        p.get("record"), p.get("supersedes_record")) if x}


def _segments(path: Path) -> tuple[list[dict], str | None]:
    """[(location, text, numbers, quoted)], owner."""
    suf = path.suffix.lower()
    segs, owner = [], None
    if suf == ".docx":
        from docx import Document
        d = Document(str(path))
        owner = d.core_properties.author or None
        for i, p in enumerate(d.paragraphs):
            if p.text.strip():
                segs.append({"loc": f"paragraph {i + 1}", "text": p.text, "numbers": [], "quoted": False})
        for ti, t in enumerate(d.tables):
            for ri, row in enumerate(t.rows):
                segs.append({"loc": f"table {ti + 1} row {ri + 1}", "text": " | ".join(c.text for c in row.cells),
                             "numbers": [], "quoted": False})
    elif suf == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(str(path), data_only=True)
        owner = wb.properties.creator or None
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                vals = [c.value for c in row if c.value is not None]
                if not vals:
                    continue
                segs.append({"loc": f"{ws.title}!row {row[0].row}", "text": " | ".join(str(v) for v in vals),
                             "numbers": [v for v in vals if isinstance(v, (int, float)) and not isinstance(v, bool)],
                             "quoted": False})
    elif suf == ".eml":
        msg = email.message_from_bytes(path.read_bytes(), policy=email.policy.default)
        owner = email.utils.parseaddr(str(msg.get("From", "")))[1] or None
        segs.append({"loc": "Subject", "text": str(msg.get("Subject", "")), "numbers": [], "quoted": False})
        for part in msg.walk():
            if part.get_content_type() == "text/plain":
                for i, line in enumerate(part.get_content().splitlines(), start=1):
                    if line.strip():
                        segs.append({"loc": f"body line {i}", "text": line, "numbers": [],
                                     "quoted": line.lstrip().startswith(">")})
    elif suf in (".md", ".txt"):
        lines = path.read_text(encoding="utf-8").splitlines()
        for i, line in enumerate(lines, start=1):
            m = re.match(r"^owner:\s*(\S+)", line.strip(), re.I)
            if m and i <= 5:
                owner = m.group(1)
            elif line.strip():
                segs.append({"loc": f"line {i}", "text": line, "numbers": [], "quoted": line.lstrip().startswith(">")})
    return segs, owner


def scan(register: list[dict], outgoing: Path, *, known_entities: list[str], responses: list[dict] | None = None,
         now: _dt.datetime | None = None) -> dict:
    guards = jsonio.read(RULES_DIR / "misc.json")["scanner_guards"]
    skip_name = re.compile(guards["skip_name_regex"])
    hist = re.compile(guards["historical_context_regex"])
    closed = {(r["entry"], r["file"]) for r in (responses or []) if r.get("already_fixed")}
    findings, suppressed, raw = [], {}, 0

    def sup(g):
        suppressed[g] = suppressed.get(g, 0) + 1

    files = sorted(p for p in Path(outgoing).rglob("*") if p.is_file() and p.suffix.lower() in
                   (".docx", ".xlsx", ".eml", ".md", ".txt"))
    cache = {}
    for e in register:
        old_rx, new_rx = form_regex(e["old"]), form_regex(e["new"])
        numeric = not _is_date(e["old"])
        old_f = float(e["old"]) if numeric else None
        new_f = float(e["new"]) if numeric else None
        cp_key = ent.match_key(e.get("counterparty") or "")
        others = [k for k in known_entities if ent.match_key(k) != cp_key]
        since = _dt.date.fromisoformat(e["since"])
        for f in files:
            rel = f.relative_to(outgoing).as_posix()
            if f not in cache:
                cache[f] = _segments(f)
            segs, owner = cache[f]
            hits = [s for s in segs if old_rx.search(s["text"]) or (numeric and any(abs(n - old_f) < 1e-9
                                                                                  for n in s["numbers"]))]
            if not hits:
                continue
            raw += len(hits)
            if skip_name.search(rel):
                sup("G1_snapshot_name")
                continue
            mtime = _dt.datetime.fromtimestamp(os.path.getmtime(f), _dt.timezone.utc).date()
            if mtime < since:
                sup("G2_untouched_since")
                continue
            doc_text = " ".join(s["text"] for s in segs)
            doc_links = (cp_key and cp_key in ent.match_key(doc_text)) or (
                e.get("pratica") and e["pratica"] in doc_text) or e["proof"]["document"] in doc_text or \
                e["proof"]["supersedes"] in doc_text
            for s in hits:
                t = s["text"]
                if new_rx.search(t) or (numeric and any(abs(n - new_f) < 1e-9 for n in s["numbers"])):
                    sup("G3_comparison")
                    continue
                if s["quoted"] or hist.search(t):
                    sup("G4_historical_context")
                    continue
                refs = set(REF_RX.findall(t))
                if refs and not (refs & own_refs(e)):
                    sup("G5_other_dossier")
                    continue
                mk = ent.match_key(t)
                names_other = any(ent.match_key(o) and ent.match_key(o) in mk for o in others)
                names_cp = bool(cp_key) and cp_key in mk
                if names_other and not names_cp:
                    sup("G5_other_counterparty")
                    continue
                if not names_cp and guards["require_counterparty_or_ref_in_document"] and not doc_links:
                    sup("G5_no_link_to_counterparty")
                    continue
                status = "closed_by_owner" if (e["id"], rel) in closed else "open"
                if status != "open":
                    sup("G6_owner_already_fixed")
                findings.append({"entry": e["id"], "file": rel, "location": s["loc"], "snippet": t[:160],
                                 "old": e["old"], "new": e["new"], "since": e["since"], "proof": e["proof"]["document"],
                                 "owner": owner or (e.get("owners") or ["amministrazione"])[0], "status": status})
    open_f = [x for x in findings if x["status"] == "open"]
    by_owner = {}
    for x in open_f:
        by_owner.setdefault(x["owner"], []).append(x)
    return {"raw_matches": raw, "suppressed": dict(sorted(suppressed.items())), "flags_open": len(open_f),
            "flags_closed_by_owner": len(findings) - len(open_f), "findings": findings,
            "by_owner": {k: by_owner[k] for k in sorted(by_owner)}}


def run(register: list[dict], outgoing_dirs: list[Path], known_entities: list[str], out_path: Path,
        responses: list[dict] | None = None) -> dict:
    results = {}
    for d in outgoing_dirs:
        if Path(d).exists():
            results[Path(d).name] = scan(register, Path(d), known_entities=known_entities, responses=responses)
    state = {"stage": "s9_distribution", "scanned": sorted(results), "results": results,
             "flags_open": sum(r["flags_open"] for r in results.values())}
    jsonio.write(out_path, state)
    return state
