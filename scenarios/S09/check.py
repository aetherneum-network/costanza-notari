"""S09 - .p7m with valid integrity, unverifiable chain (Officine Lagorai S.r.l.). Pass: signer_chain_verified: false shown (L8)."""
import shutil
import subprocess
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent.parent))
from scenarios._common import ROOT, ensure_corpus, load, main, report, run_pipeline, workdir  # noqa: E402


def _openssl():
    for c in (shutil.which("openssl"), r"C:\Program Files\Git\mingw64\bin\openssl.exe"):
        if c and Path(c).exists():
            return c
    return None


def check():
    ensure_corpus()
    w = workdir("S09")
    code, _ = run_pipeline(HERE / "input" / "corpus", w, config=HERE / "input" / "config.json",
                           as_of="2026-10-21T09:40:00+02:00")
    sig = load(w / "work" / "state" / "03_signatures.json")["records"][0]["attachments"][0]
    exp = load(HERE / "expected" / "expected.json")
    fields_ok = all(sig[k] == v for k, v in exp["signature"].items())
    from openpyxl import load_workbook
    wb = load_workbook(w / "work" / "build" / "master_index.xlsx")
    idx = [c.value for c in wb["Index"][7]]
    sig_row = [c.value for c in wb["Signatures"][2]]
    shown = idx[15] == exp["index_chain_column"] and sig_row[3:6] == exp["signatures_sheet"]
    from docx import Document
    report_text = "\n".join(c.text for t in Document(str(w / "work" / "build" / "report.docx")).tables for r in t.rows
                            for c in r.cells)
    in_report = exp["report_row"] in report_text
    ossl, cross = _openssl(), "n/a"
    if ossl:
        p7m = next((w / "work" / "att" / "ARC-0001").glob("*.p7m"))
        a = subprocess.run([ossl, "cms", "-verify", "-noverify", "-binary", "-inform", "DER", "-in", str(p7m),
                            "-out", str(w / "x.pdf")], capture_output=True).returncode
        b = subprocess.run([ossl, "cms", "-verify", "-binary", "-inform", "DER", "-in", str(p7m), "-CAfile",
                            str(ROOT / "corpus" / "out" / "testca" / "trust" / "test-root-ca.pem"), "-purpose", "any",
                            "-out", str(w / "y.pdf")], capture_output=True).returncode
        cross = f"openssl -noverify rc={a} (integrity+extract), with -CAfile rc={b} (chain)"
        shown = shown and a == 0 and b != 0
    ok = code == 0 and fields_ok and shown and in_report
    return report("S09", ok, f"signature_integrity={sig['signature_integrity']}, signer_chain_verified="
                  f"{str(sig['signer_chain_verified']).lower()} ({sig['chain_status']}, issuer '{sig['signer']['issuer_cn']}'); "
                  f"index column 'Signer chain verified' = {idx[15]}; Signatures sheet {sig_row[3:6]}; in report: {in_report}; "
                  f"{cross}", {"signature": sig})


if __name__ == "__main__":
    main(check)
