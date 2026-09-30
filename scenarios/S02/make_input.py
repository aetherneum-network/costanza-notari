"""Generate S02 input (deterministic): two editions of an instalment notice and the outgoing
documents of the synthetic debtor Cantine Belvedere S.p.A. - 5 planted stale values, 10 decoys,
1 stale value already fixed by its owner."""
import datetime as dt
import json
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent.parent))
from pipeline.lib import zipnorm  # noqa: E402

IN = HERE / "input"
OUT = IN / "outgoing"
T0 = dt.datetime(2026, 10, 21, 7, 0, tzinfo=dt.timezone.utc)


def docx(rel, author, paragraphs, table=None):
    from docx import Document
    d = Document()
    for p in paragraphs:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=0, cols=len(table[0]))
        for row in table:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = v
    cp = d.core_properties
    cp.author, cp.last_modified_by, cp.title = author, author, rel
    cp.created = cp.modified = T0.replace(tzinfo=None)
    cp.revision = 1
    p = OUT / rel
    p.parent.mkdir(parents=True, exist_ok=True)
    d.save(p)
    zipnorm.normalize(p, T0)


def xlsx(rel, creator, rows):
    from openpyxl import Workbook
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    wb.properties.creator = creator
    wb.properties.created = wb.properties.modified = T0.replace(tzinfo=None)
    p = OUT / rel
    p.parent.mkdir(parents=True, exist_ok=True)
    wb.save(p)
    zipnorm.normalize(p, T0)


def eml(rel, sender, subject, body):
    p = OUT / rel
    p.parent.mkdir(parents=True, exist_ok=True)
    raw = (f"From: {sender}\r\nTo: direzione@cantinebelvedere.example\r\nSubject: {subject}\r\n"
           "Date: Wed, 21 Oct 2026 09:00:00 +0200\r\nMIME-Version: 1.0\r\n"
           "Content-Type: text/plain; charset=UTF-8\r\nContent-Transfer-Encoding: 8bit\r\n\r\n"
           + body.replace("\n", "\r\n") + "\r\n")
    p.write_bytes(raw.encode("utf-8"))


def md(rel, owner, lines):
    p = OUT / rel
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(f"owner: {owner}\n\n" + "\n".join(lines) + "\n", encoding="utf-8", newline="\n")


def main():
    base = {"doc_type": "provvedimento_rateizzazione", "area": "collection", "party_entity": "AGENZIA ESEMPIO RISCOSSIONE",
            "pratica": "PR-RAT-0077", "sender_class": "CORPORATE_PEC", "transmitter_entity": "AGENZIA ESEMPIO RISCOSSIONE",
            "author_entity": "AGENZIA ESEMPIO RISCOSSIONE", "deadlines": [{"date": "2026-11-30", "nature": "actionable"}],
            "deadline": "2026-11-30", "deadline_nature": "actionable", "recuperare_fields": []}
    notices = [
        {**base, "record_id": "ARC-0101", "envelope": "envelopes/2026-10/PEC_20261002_101500_0101.eml",
         "content_sha256": "11" * 32, "edition_ref": "EX-2210", "edition_date": "2026-10-02", "amount_due": "3415.20",
         "pec_time": "2026-10-02T10:15:00+02:00", "notification_date": "2026-10-02"},
        {**base, "record_id": "ARC-0102", "envelope": "envelopes/2026-10/PEC_20261020_093000_0102.eml",
         "content_sha256": "22" * 32, "edition_ref": "EX-2231", "edition_date": "2026-10-20", "amount_due": "3154.20",
         "supersedes_ref": "EX-2210", "pec_time": "2026-10-20T09:30:00+02:00", "notification_date": "2026-10-20"},
    ]
    IN.mkdir(exist_ok=True)
    (IN / "notices.json").write_text(json.dumps({"records": notices}, indent=2) + "\n", encoding="utf-8", newline="\n")
    dates = {}
    # --- planted stale values (5)
    docx("lettere/lettera_accompagnamento_2026-10-21.docx", "amministrazione",
         ["Spett.le AGENZIA ESEMPIO RISCOSSIONE", "Oggetto: piano di rateizzazione PR-RAT-0077",
          "Come concordato, la rata mensile del piano è pari a € 3.415,20 e sarà versata entro il 30/11/2026."])
    dates["lettere/lettera_accompagnamento_2026-10-21.docx"] = "2026-10-21"
    eml("email/riepilogo_2026-10-21.eml", "tesoreria@cantinebelvedere.example", "Riepilogo scadenze",
        "Buongiorno,\nRata piano PR-RAT-0077 (Agenzia Esempio Riscossione): 3.415,20 euro.\nSaluti")
    dates["email/riepilogo_2026-10-21.eml"] = "2026-10-21"
    xlsx("tesoreria/previsione_cassa_Q4.xlsx", "tesoreria",
         [["Controparte", "Pratica", "Rata"], ["AGENZIA ESEMPIO RISCOSSIONE", "PR-RAT-0077", 3415.2]])
    dates["tesoreria/previsione_cassa_Q4.xlsx"] = "2026-10-22"
    md("note/appunti_tesoreria.md", "tesoreria", ["- PR-RAT-0077: pagare 3,415.20 il 30/11"])
    dates["note/appunti_tesoreria.md"] = "2026-10-23"
    docx("direzione/memo_cda.docx", "direzione", ["Memo per il CdA - impegni di pagamento"],
         table=[["Controparte", "Voce", "Importo"], ["AGENZIA ESEMPIO RISCOSSIONE", "rata mensile", "3.415,20"]])
    dates["direzione/memo_cda.docx"] = "2026-10-21"
    # --- stale value already fixed by its owner (closed via owner_responses.json)
    md("lettere/bozza_sollecito.md", "amministrazione", ["Rata PR-RAT-0077: € 3.415,20"])
    dates["lettere/bozza_sollecito.md"] = "2026-10-22"
    # --- decoys (10)
    docx("lettere/lettera_accompagnamento_2026-10-15.docx", "amministrazione",
         ["Oggetto: piano di rateizzazione PR-RAT-0077", "La rata mensile del piano è pari a € 3.415,20."])
    dates["lettere/lettera_accompagnamento_2026-10-15.docx"] = "2026-10-15"          # D1 untouched since
    xlsx("tesoreria/snapshot_riepilogo_2026-10-10.xlsx", "tesoreria",
         [["AGENZIA ESEMPIO RISCOSSIONE", "PR-RAT-0077", 3415.2]])
    dates["tesoreria/snapshot_riepilogo_2026-10-10.xlsx"] = "2026-10-22"            # D2 snapshot by name
    md("archivio/lettera_2026-10-03.md", "amministrazione", ["Rata PR-RAT-0077: € 3.415,20"])
    dates["archivio/lettera_2026-10-03.md"] = "2026-10-22"                           # D3 archive by name
    eml("email/risposta_con_citazione.eml", "direzione@cantinebelvedere.example", "Re: rata",
        "Ricevuto, grazie.\n> la rata del piano PR-RAT-0077 è di 3.415,20 euro")
    dates["email/risposta_con_citazione.eml"] = "2026-10-21"                         # D4 quoted reply
    md("note/confronto.md", "tesoreria", ["Rata PR-RAT-0077 ridotta da € 3.415,20 a € 3.154,20."])
    dates["note/confronto.md"] = "2026-10-21"                                        # D5 comparison
    md("note/riepilogo_versamenti.md", "tesoreria", ["Totale versato finora sul piano PR-RAT-0077: € 3.415,20."])
    dates["note/riepilogo_versamenti.md"] = "2026-10-21"                                          # D6 historical total
    md("note/altro_fornitore.md", "tesoreria", ["TESSITURE MONTEVERDE S.R.L. - fattura saldata € 3.415,20."])
    dates["note/altro_fornitore.md"] = "2026-10-21"                                  # D7 other counterparty
    xlsx("tesoreria/scadenziario_fornitori.xlsx", "tesoreria",
         [["Controparte", "Pratica", "Importo"], ["TESSITURE MONTEVERDE S.R.L.", "PR-TM-0001", 3415.2]])
    dates["tesoreria/scadenziario_fornitori.xlsx"] = "2026-10-21"                    # D8 other dossier
    xlsx("tesoreria/importi_vicini.xlsx", "tesoreria",
         [["AGENZIA ESEMPIO RISCOSSIONE", "PR-RAT-0077", 3415.25, "13.415,20", 34152]])
    dates["tesoreria/importi_vicini.xlsx"] = "2026-10-21"                            # D9 near numbers
    md("note/preventivo_bottiglie.md", "acquisti", ["Preventivo vetreria: € 3.415,20 per 12.000 bottiglie."])
    dates["note/preventivo_bottiglie.md"] = "2026-10-21"                             # D10 same number, other sense
    (IN / "file_dates.json").write_text(json.dumps(dict(sorted(dates.items())), indent=2) + "\n", encoding="utf-8",
                                        newline="\n")
    (IN / "owner_responses.json").write_text(json.dumps([{"entry": "SV-0002-amount_due", "file": "lettere/bozza_sollecito.md",
                                                          "already_fixed": True, "by": "amministrazione",
                                                          "note": "fixed in the version that was sent"}], indent=2)
                                             + "\n", encoding="utf-8", newline="\n")


if __name__ == "__main__":
    main()
